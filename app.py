import io
import os
import sqlite3
from datetime import datetime, timedelta
from functools import wraps

from flask import (
    Flask, request, session, redirect, url_for, render_template,
    flash, g, abort, send_file
)
from werkzeug.security import generate_password_hash, check_password_hash

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

APP_SECRET_KEY = os.environ.get("APP_SECRET_KEY", "dev-secret-change-me")
ADMIN_USER = os.environ.get("ADMIN_USER", "admin")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "cambiar-esta-clave")
DB_PATH = os.environ.get("DB_PATH", os.path.join(os.path.dirname(__file__), "inventario.db"))

app = Flask(__name__)
app.secret_key = APP_SECRET_KEY

AREAS = {
    "pyp": "Mantenimiento PyP",
    "mejoras": "Mantenimiento Mejoras",
}

# Qué área(s) puede operar cada rol. Un rol con más de un área se comporta
# como "multi-área": ve/filtra entre ambas en vez de estar fijo a una.
ROLE_AREAS = {
    "admin": ["pyp", "mejoras"],
    "planificador": ["pyp", "mejoras"],
    "preparador_kit": ["pyp"],
    "electromecanico": ["mejoras"],
}

ROLE_LABELS = {
    "admin": "Administrador",
    "planificador": "Planificador",
    "preparador_kit": "Preparador de Kit",
    "electromecanico": "Electromecánico",
}

ROLES_VALIDOS = tuple(ROLE_AREAS.keys())

# Códigos de persona / turno 7x7 que ingresan stock. Lista preliminar,
# ampliable sin tocar la base de datos (es solo texto libre validado aquí).
TURNOS = ["38", "44"]

# Umbrales para clasificar el stock de un componente frente a su mínimo.
UMBRAL_PUNTO_PEDIDO = 1.5  # <= 1.5x el mínimo => "a punto de estar crítico"


# ---------------------------------------------------------------------------
# Base de datos
# ---------------------------------------------------------------------------

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = sqlite3.connect(DB_PATH)
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            nombre TEXT NOT NULL,
            rol TEXT NOT NULL CHECK (rol IN ('admin', 'planificador', 'preparador_kit', 'electromecanico')),
            activo INTEGER NOT NULL DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS componentes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo TEXT UNIQUE NOT NULL,
            nombre TEXT NOT NULL,
            descripcion TEXT,
            categoria TEXT,
            ubicacion TEXT,
            area TEXT NOT NULL CHECK (area IN ('pyp', 'mejoras')),
            cantidad INTEGER NOT NULL DEFAULT 0,
            stock_minimo INTEGER NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS movimientos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            componente_id INTEGER,
            codigo TEXT NOT NULL,
            nombre TEXT NOT NULL,
            area TEXT NOT NULL,
            tipo TEXT NOT NULL CHECK (tipo IN ('ingreso', 'retiro')),
            cantidad INTEGER NOT NULL,
            usuario TEXT NOT NULL,
            turno TEXT,
            fecha TEXT NOT NULL,
            FOREIGN KEY (componente_id) REFERENCES componentes(id)
        );
        """
    )
    # Usuario admin por defecto (idempotente)
    db.execute(
        "INSERT OR IGNORE INTO usuarios (username, password_hash, nombre, rol) "
        "VALUES (?, ?, ?, 'admin')",
        (ADMIN_USER, generate_password_hash(ADMIN_PASSWORD), "Administrador"),
    )
    db.commit()
    db.close()


init_db()


# ---------------------------------------------------------------------------
# Auth y control de acceso
# ---------------------------------------------------------------------------

def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)
    return wrapped


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if session.get("rol") != "admin":
            abort(403)
        return view(*args, **kwargs)
    return wrapped


def areas_del_usuario():
    """Lista de áreas que puede tocar el rol actual (vacía si el rol no existe)."""
    return ROLE_AREAS.get(session.get("rol"), [])


def current_user_area():
    """Área de trabajo efectiva de la sesión actual.

    - Rol de una sola área (preparador_kit, electromecanico): siempre esa área.
    - Rol multi-área (admin, planificador): el área elegida en session['area_activa'],
      o None si todavía no ha elegido con qué área trabajar.
    """
    areas = areas_del_usuario()
    if len(areas) == 1:
        return areas[0]
    if len(areas) > 1:
        activa = session.get("area_activa")
        return activa if activa in areas else None
    return None


def area_permitida(area):
    """¿El usuario actual puede operar sobre esta área?"""
    return area in areas_del_usuario()


def area_requerida(view):
    """Exige que ya haya un área activa elegida (roles multi-área deben elegir primero)."""
    @wraps(view)
    def wrapped(*args, **kwargs):
        if current_user_area() is None:
            return redirect(url_for("inicio"))
        return view(*args, **kwargs)
    return wrapped


@app.context_processor
def inject_globals():
    areas = areas_del_usuario()
    multi_area = len(areas) > 1
    area_activa = current_user_area()
    return {
        "AREAS": AREAS,
        "ROLE_LABELS": ROLE_LABELS,
        "ROLE_AREAS": ROLE_AREAS,
        "TURNOS": TURNOS,
        "session_nombre": session.get("nombre"),
        "session_rol": session.get("rol"),
        "session_user_id": session.get("user_id"),
        "session_area_nombre": AREAS.get(area_activa, ""),
        "session_area_activa": area_activa,
        "session_multi_area": multi_area,
    }


# ---------------------------------------------------------------------------
# Login / logout
# ---------------------------------------------------------------------------

@app.route("/", methods=["GET"])
def index():
    return redirect(url_for("inicio") if "user_id" in session else url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        db = get_db()
        user = db.execute(
            "SELECT * FROM usuarios WHERE username = ? AND activo = 1", (username,)
        ).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["nombre"] = user["nombre"]
            session["rol"] = user["rol"]
            next_url = request.args.get("next") or url_for("inicio")
            return redirect(next_url)
        flash("Usuario o contraseña incorrectos.", "error")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ---------------------------------------------------------------------------
# Inicio / Dashboard
# ---------------------------------------------------------------------------

@app.route("/inicio")
@login_required
def inicio():
    if len(areas_del_usuario()) > 1 and current_user_area() is None:
        return render_template("elegir_area.html")
    return render_template("inicio.html")


@app.route("/area/<area>")
@login_required
def elegir_area(area):
    if area not in areas_del_usuario():
        abort(403)
    session["area_activa"] = area
    return redirect(url_for("inicio"))


@app.route("/area/cambiar")
@login_required
def cambiar_area():
    session.pop("area_activa", None)
    return redirect(url_for("inicio"))


@app.route("/dashboard")
@login_required
@area_requerida
def dashboard():
    db = get_db()
    area = current_user_area()

    total_componentes = db.execute(
        "SELECT COUNT(*) FROM componentes WHERE area = ?", (area,)
    ).fetchone()[0]
    total_unidades = db.execute(
        "SELECT COALESCE(SUM(cantidad), 0) FROM componentes WHERE area = ?", (area,)
    ).fetchone()[0]
    sin_stock = db.execute(
        "SELECT COUNT(*) FROM componentes WHERE area = ? AND cantidad <= stock_minimo", (area,)
    ).fetchone()[0]

    hoy = datetime.now().strftime("%Y-%m-%d")
    movimientos_hoy = db.execute(
        "SELECT COUNT(*) FROM movimientos WHERE fecha LIKE ? AND area = ?",
        (f"{hoy}%", area),
    ).fetchone()[0]

    ultimos_movimientos = db.execute(
        "SELECT * FROM movimientos WHERE area = ? ORDER BY id DESC LIMIT 10", (area,)
    ).fetchall()

    bajo_stock = db.execute(
        "SELECT * FROM componentes WHERE area = ? AND cantidad <= stock_minimo "
        "ORDER BY nombre LIMIT 10",
        (area,),
    ).fetchall()

    return render_template(
        "dashboard.html",
        total_componentes=total_componentes,
        total_unidades=total_unidades,
        sin_stock=sin_stock,
        movimientos_hoy=movimientos_hoy,
        ultimos_movimientos=ultimos_movimientos,
        bajo_stock=bajo_stock,
    )


# ---------------------------------------------------------------------------
# Componentes (listado, alta, edición, baja)
# ---------------------------------------------------------------------------

@app.route("/componentes")
@login_required
@area_requerida
def componentes_list():
    db = get_db()
    area = current_user_area()
    q = request.args.get("q", "").strip()

    sql = "SELECT * FROM componentes WHERE area = ?"
    params = [area]
    if q:
        sql += " AND (codigo LIKE ? OR nombre LIKE ? OR categoria LIKE ? OR ubicacion LIKE ?)"
        like = f"%{q}%"
        params.extend([like, like, like, like])
    sql += " ORDER BY nombre"

    componentes = db.execute(sql, params).fetchall()
    return render_template(
        "componentes_list.html",
        componentes=componentes,
        q=q,
    )


@app.route("/componentes/nuevo", methods=["GET", "POST"])
@login_required
@area_requerida
def componente_nuevo():
    area = current_user_area()
    if request.method == "POST":
        codigo = request.form.get("codigo", "").strip()
        nombre = request.form.get("nombre", "").strip()
        descripcion = request.form.get("descripcion", "").strip()
        categoria = request.form.get("categoria", "").strip()
        ubicacion = request.form.get("ubicacion", "").strip()
        cantidad = request.form.get("cantidad", "0")
        stock_minimo = request.form.get("stock_minimo", "0")

        if not codigo or not nombre:
            flash("Código SAP y nombre son obligatorios.", "error")
            return render_template("componente_form.html", componente=request.form)

        db = get_db()
        try:
            db.execute(
                "INSERT INTO componentes (codigo, nombre, descripcion, categoria, ubicacion, "
                "area, cantidad, stock_minimo) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (codigo, nombre, descripcion, categoria, ubicacion, area,
                 int(cantidad or 0), int(stock_minimo or 0)),
            )
            db.commit()
        except sqlite3.IntegrityError:
            flash(f'Ya existe un componente con el código SAP "{codigo}".', "error")
            return render_template("componente_form.html", componente=request.form)

        flash("Componente creado correctamente.", "success")
        return redirect(url_for("componentes_list"))

    return render_template("componente_form.html", componente=None)


@app.route("/componentes/<int:comp_id>/editar", methods=["GET", "POST"])
@login_required
def componente_editar(comp_id):
    db = get_db()
    componente = db.execute("SELECT * FROM componentes WHERE id = ?", (comp_id,)).fetchone()
    if componente is None:
        abort(404)
    if not area_permitida(componente["area"]):
        abort(403)

    if request.method == "POST":
        nombre = request.form.get("nombre", "").strip()
        descripcion = request.form.get("descripcion", "").strip()
        categoria = request.form.get("categoria", "").strip()
        ubicacion = request.form.get("ubicacion", "").strip()
        stock_minimo = request.form.get("stock_minimo", "0")

        if not nombre:
            flash("El nombre es obligatorio.", "error")
            return render_template("componente_form.html", componente=componente, editar=True)

        db.execute(
            "UPDATE componentes SET nombre = ?, descripcion = ?, categoria = ?, "
            "ubicacion = ?, stock_minimo = ? WHERE id = ?",
            (nombre, descripcion, categoria, ubicacion, int(stock_minimo or 0), comp_id),
        )
        db.commit()
        flash("Componente actualizado.", "success")
        return redirect(url_for("componentes_list"))

    return render_template("componente_form.html", componente=componente, editar=True)


@app.route("/componentes/<int:comp_id>/eliminar", methods=["POST"])
@login_required
def componente_eliminar(comp_id):
    db = get_db()
    componente = db.execute("SELECT * FROM componentes WHERE id = ?", (comp_id,)).fetchone()
    if componente is None:
        abort(404)
    if not area_permitida(componente["area"]):
        abort(403)

    db.execute("DELETE FROM componentes WHERE id = ?", (comp_id,))
    db.commit()
    flash("Componente eliminado.", "success")
    return redirect(url_for("componentes_list"))


# ---------------------------------------------------------------------------
# Ingreso / Retiro de stock
# ---------------------------------------------------------------------------

def _componentes_disponibles(area):
    db = get_db()
    if area:
        return db.execute(
            "SELECT * FROM componentes WHERE area = ? ORDER BY nombre", (area,)
        ).fetchall()
    return db.execute("SELECT * FROM componentes ORDER BY area, nombre").fetchall()


def _fecha_valida_o_hoy(valor):
    """Valida un input type=date (YYYY-MM-DD); si falta o es inválido, usa hoy."""
    hoy = datetime.now().strftime("%Y-%m-%d")
    if not valor:
        return hoy
    try:
        datetime.strptime(valor, "%Y-%m-%d")
        return valor
    except ValueError:
        return hoy


@app.route("/ingreso", methods=["GET", "POST"])
@login_required
@area_requerida
def ingreso():
    area_activa = current_user_area()
    db = get_db()

    if request.method == "POST":
        modo = request.form.get("modo", "existente")
        cantidad = int(request.form.get("cantidad", "0") or 0)

        turno = request.form.get("turno", "").strip()
        if request.form.get("turno") == "otro":
            turno = request.form.get("turno_otro", "").strip()
        fecha_elegida = _fecha_valida_o_hoy(request.form.get("fecha", "").strip())

        if cantidad <= 0:
            flash("La cantidad debe ser mayor a 0.", "error")
            return redirect(url_for("ingreso"))
        if not turno:
            flash("El código de turno de la persona que ingresa es obligatorio.", "error")
            return redirect(url_for("ingreso"))

        if modo == "nuevo":
            codigo = request.form.get("codigo", "").strip()
            nombre = request.form.get("nombre", "").strip()
            stock_minimo = request.form.get("stock_minimo", "0")
            if not codigo or not nombre:
                flash("Código SAP y nombre son obligatorios para un componente nuevo.", "error")
                return redirect(url_for("ingreso"))
            categoria = request.form.get("categoria", "").strip()
            ubicacion = request.form.get("ubicacion", "").strip()
            try:
                cur = db.execute(
                    "INSERT INTO componentes (codigo, nombre, categoria, ubicacion, area, "
                    "cantidad, stock_minimo) VALUES (?, ?, ?, ?, ?, ?, ?)",
                    (codigo, nombre, categoria, ubicacion, area_activa, cantidad, int(stock_minimo or 0)),
                )
                componente_id = cur.lastrowid
                area = area_activa
            except sqlite3.IntegrityError:
                flash(f'Ya existe un componente con el código SAP "{codigo}". Usa "Ingreso a existente".', "error")
                return redirect(url_for("ingreso"))
        else:
            componente_id = request.form.get("componente_id")
            componente = db.execute(
                "SELECT * FROM componentes WHERE id = ?", (componente_id,)
            ).fetchone()
            if componente is None or componente["area"] != area_activa:
                flash("Selecciona un componente válido.", "error")
                return redirect(url_for("ingreso"))
            db.execute(
                "UPDATE componentes SET cantidad = cantidad + ? WHERE id = ?",
                (cantidad, componente_id),
            )
            codigo, nombre, area = componente["codigo"], componente["nombre"], componente["area"]

        hora_actual = datetime.now().strftime("%H:%M:%S")
        db.execute(
            "INSERT INTO movimientos (componente_id, codigo, nombre, area, tipo, cantidad, "
            "usuario, turno, fecha) VALUES (?, ?, ?, ?, 'ingreso', ?, ?, ?, ?)",
            (componente_id, codigo, nombre, area, cantidad, session["username"], turno,
             f"{fecha_elegida} {hora_actual}"),
        )
        db.commit()
        flash(f"Ingreso registrado: +{cantidad} de {nombre}.", "success")
        return redirect(url_for("ingreso"))

    componentes = _componentes_disponibles(area_activa)
    hoy = datetime.now().strftime("%Y-%m-%d")
    return render_template("ingreso.html", componentes=componentes, hoy=hoy)


@app.route("/retiro", methods=["GET", "POST"])
@login_required
@area_requerida
def retiro():
    area_activa = current_user_area()
    db = get_db()

    if request.method == "POST":
        componente_id = request.form.get("componente_id")
        cantidad = int(request.form.get("cantidad", "0") or 0)
        componente = db.execute(
            "SELECT * FROM componentes WHERE id = ?", (componente_id,)
        ).fetchone()

        if componente is None or componente["area"] != area_activa:
            flash("Selecciona un componente válido.", "error")
        elif cantidad <= 0:
            flash("La cantidad debe ser mayor a 0.", "error")
        elif cantidad > componente["cantidad"]:
            flash(
                f'No hay stock suficiente de "{componente["nombre"]}" '
                f'(disponible: {componente["cantidad"]}).', "error"
            )
        else:
            db.execute(
                "UPDATE componentes SET cantidad = cantidad - ? WHERE id = ?",
                (cantidad, componente_id),
            )
            db.execute(
                "INSERT INTO movimientos (componente_id, codigo, nombre, area, tipo, cantidad, "
                "usuario, fecha) VALUES (?, ?, ?, ?, 'retiro', ?, ?, ?)",
                (componente_id, componente["codigo"], componente["nombre"], componente["area"],
                 cantidad, session["username"], datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            )
            db.commit()
            flash(f'Retiro registrado: -{cantidad} de {componente["nombre"]}.', "success")
        return redirect(url_for("retiro"))

    componentes = [c for c in _componentes_disponibles(area_activa) if c["cantidad"] > 0]
    return render_template("retiro.html", componentes=componentes)


# ---------------------------------------------------------------------------
# Búsqueda
# ---------------------------------------------------------------------------

@app.route("/buscar")
@login_required
@area_requerida
def buscar():
    area = current_user_area()
    q = request.args.get("q", "").strip()
    resultados = []
    if q:
        db = get_db()
        sql = ("SELECT * FROM componentes WHERE area = ? AND "
               "(codigo LIKE ? OR nombre LIKE ? OR categoria LIKE ? OR ubicacion LIKE ?)")
        like = f"%{q}%"
        params = [area, like, like, like, like]
        sql += " ORDER BY nombre"
        resultados = db.execute(sql, params).fetchall()
    return render_template("buscar.html", q=q, resultados=resultados)


# ---------------------------------------------------------------------------
# Tendencias de uso (para ajustar puntos de pedido)
# ---------------------------------------------------------------------------

@app.route("/tendencias")
@login_required
@area_requerida
def tendencias():
    db = get_db()
    area = current_user_area()

    dias = int(request.args.get("dias", "90") or 90)
    desde = (datetime.now() - timedelta(days=dias)).strftime("%Y-%m-%d")

    mas_usados = db.execute(
        """
        SELECT m.codigo, m.nombre, m.area, SUM(m.cantidad) AS total_retirado,
               COUNT(*) AS movimientos
        FROM movimientos m
        WHERE m.tipo = 'retiro' AND m.fecha >= ? AND m.area = ?
        GROUP BY m.componente_id, m.codigo, m.nombre, m.area
        ORDER BY total_retirado DESC
        LIMIT 15
        """,
        (desde, area),
    ).fetchall()

    componentes = db.execute(
        "SELECT * FROM componentes WHERE area = ?", (area,)
    ).fetchall()

    usados_ids = set()
    uso_por_componente = {}
    for row in db.execute(
        """
        SELECT componente_id, SUM(cantidad) AS total
        FROM movimientos m
        WHERE tipo = 'retiro' AND fecha >= ? AND m.area = ?
        GROUP BY componente_id
        """,
        (desde, area),
    ).fetchall():
        if row["componente_id"] is not None:
            usados_ids.add(row["componente_id"])
            uso_por_componente[row["componente_id"]] = row["total"]

    sin_movimiento = [c for c in componentes if c["id"] not in usados_ids]

    sugerencias = []
    semanas = max(dias / 7.0, 1)
    for c in componentes:
        total = uso_por_componente.get(c["id"], 0)
        promedio_semanal = total / semanas
        if total > 0 and promedio_semanal * 2 > c["stock_minimo"]:
            sugerencias.append({
                "componente": c,
                "promedio_semanal": round(promedio_semanal, 1),
                "tipo": "subir",
            })
        elif total == 0 and c["stock_minimo"] > 0:
            sugerencias.append({
                "componente": c,
                "promedio_semanal": 0,
                "tipo": "revisar",
            })

    return render_template(
        "tendencias.html",
        dias=dias,
        mas_usados=mas_usados,
        sin_movimiento=sin_movimiento[:20],
        sugerencias=sugerencias[:20],
    )


# ---------------------------------------------------------------------------
# Informe de stock en Excel
# ---------------------------------------------------------------------------

def _clasificar_stock(cantidad, stock_minimo):
    if cantidad <= 0:
        return "Crítico"
    if cantidad <= stock_minimo:
        return "Bajo stock"
    if stock_minimo > 0 and cantidad <= stock_minimo * UMBRAL_PUNTO_PEDIDO:
        return "A punto de estar crítico"
    return "Disponible"


@app.route("/reporte/stock.xlsx")
@login_required
@area_requerida
def reporte_stock_excel():
    db = get_db()
    area = current_user_area()

    componentes = db.execute(
        "SELECT * FROM componentes WHERE area = ? ORDER BY nombre", (area,)
    ).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"

    headers = ["Código SAP", "Nombre", "Categoría", "Ubicación",
               "Cantidad", "Stock mínimo", "Estado"]
    ws.append(headers)
    header_fill = PatternFill(start_color="0F2540", end_color="0F2540", fill_type="solid")
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    estado_fill = {
        "Crítico": PatternFill(start_color="F8D0CB", end_color="F8D0CB", fill_type="solid"),
        "Bajo stock": PatternFill(start_color="FCE8CC", end_color="FCE8CC", fill_type="solid"),
        "A punto de estar crítico": PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid"),
        "Disponible": PatternFill(start_color="D9F2E3", end_color="D9F2E3", fill_type="solid"),
    }

    resumen = {"Crítico": 0, "Bajo stock": 0, "A punto de estar crítico": 0, "Disponible": 0}
    for c in componentes:
        estado = _clasificar_stock(c["cantidad"], c["stock_minimo"])
        resumen[estado] += 1
        row = [
            c["codigo"], c["nombre"], c["categoria"] or "", c["ubicacion"] or "",
            c["cantidad"], c["stock_minimo"], estado,
        ]
        ws.append(row)
        estado_cell = ws.cell(row=ws.max_row, column=7)
        estado_cell.fill = estado_fill.get(estado)

    for col_idx, header in enumerate(headers, start=1):
        width = max(len(header) + 4, 14)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    resumen_ws = wb.create_sheet("Resumen")
    resumen_ws.append(["Área", AREAS.get(area, area)])
    resumen_ws.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M")])
    resumen_ws.append([])
    resumen_ws.append(["Estado", "Cantidad de componentes"])
    for estado, cantidad in resumen.items():
        resumen_ws.append([estado, cantidad])
    resumen_ws.column_dimensions["A"].width = 26
    resumen_ws.column_dimensions["B"].width = 22

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre_archivo = f"informe_stock_{area}_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nombre_archivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---------------------------------------------------------------------------
# Administración de usuarios (solo admin)
# ---------------------------------------------------------------------------

@app.route("/admin/usuarios")
@login_required
@admin_required
def usuarios_list():
    db = get_db()
    usuarios = db.execute("SELECT * FROM usuarios ORDER BY rol, username").fetchall()
    return render_template("usuarios.html", usuarios=usuarios)


@app.route("/admin/usuarios/nuevo", methods=["GET", "POST"])
@login_required
@admin_required
def usuario_nuevo():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        nombre = request.form.get("nombre", "").strip()
        rol = request.form.get("rol", "")
        password = request.form.get("password", "")

        if not username or not nombre or rol not in ROLES_VALIDOS or not password:
            flash("Todos los campos son obligatorios.", "error")
            return render_template("usuario_form.html", usuario=request.form)

        db = get_db()
        try:
            db.execute(
                "INSERT INTO usuarios (username, password_hash, nombre, rol) VALUES (?, ?, ?, ?)",
                (username, generate_password_hash(password), nombre, rol),
            )
            db.commit()
        except sqlite3.IntegrityError:
            flash(f'Ya existe un usuario con el nombre de usuario "{username}".', "error")
            return render_template("usuario_form.html", usuario=request.form)

        flash("Usuario creado correctamente.", "success")
        return redirect(url_for("usuarios_list"))

    return render_template("usuario_form.html", usuario=None)


@app.route("/admin/usuarios/<int:user_id>/editar", methods=["GET", "POST"])
@login_required
@admin_required
def usuario_editar(user_id):
    db = get_db()
    usuario = db.execute("SELECT * FROM usuarios WHERE id = ?", (user_id,)).fetchone()
    if usuario is None:
        abort(404)

    if request.method == "POST":
        nombre = request.form.get("nombre", "").strip()
        rol = request.form.get("rol", "")
        password = request.form.get("password", "").strip()
        activo = 1 if request.form.get("activo") == "on" else 0

        if not nombre or rol not in ROLES_VALIDOS:
            flash("Nombre y rol son obligatorios.", "error")
            return render_template("usuario_form.html", usuario=usuario, editar=True)

        if password:
            db.execute(
                "UPDATE usuarios SET nombre = ?, rol = ?, password_hash = ?, activo = ? WHERE id = ?",
                (nombre, rol, generate_password_hash(password), activo, user_id),
            )
        else:
            db.execute(
                "UPDATE usuarios SET nombre = ?, rol = ?, activo = ? WHERE id = ?",
                (nombre, rol, activo, user_id),
            )
        db.commit()
        flash("Usuario actualizado.", "success")
        return redirect(url_for("usuarios_list"))

    return render_template("usuario_form.html", usuario=usuario, editar=True)


@app.route("/admin/usuarios/<int:user_id>/eliminar", methods=["POST"])
@login_required
@admin_required
def usuario_eliminar(user_id):
    if user_id == session.get("user_id"):
        flash("No puedes eliminar tu propio usuario mientras tienes la sesión abierta.", "error")
        return redirect(url_for("usuarios_list"))
    db = get_db()
    db.execute("DELETE FROM usuarios WHERE id = ?", (user_id,))
    db.commit()
    flash("Usuario eliminado.", "success")
    return redirect(url_for("usuarios_list"))


@app.errorhandler(403)
def forbidden(e):
    return render_template("error.html", codigo=403,
                            mensaje="No tienes acceso a esta área o sección."), 403


@app.errorhandler(404)
def not_found(e):
    return render_template("error.html", codigo=404,
                            mensaje="No se encontró lo que buscabas."), 404


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
